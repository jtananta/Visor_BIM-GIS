"""Convierte el Excel de control de paquetes y componentes en data/componentes.csv.

Uso:
    python catalogo.py "PAQUETES Y COMPONENTES.xlsx" <carpeta_salida>

Lee la hoja "CONTROL GENERAL" (SECTOR, N° PAQUETE, COMPONENTE, DESCRIPCIÓN, ESTADO, % AVANCE,
RESPONSABLE, FECHA ACTUALIZACIÓN, OBSERVACIONES). El estado que el visor muestra en la ficha sale
de aquí: basta con actualizar el Excel y volver a correr este script, sin reconvertir los IFC.
"""

import csv
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

COLUMNAS = {
    "SECTOR": "sector", "N° PAQUETE": "paquete", "COMPONENTE": "componente", "DESCRIPCIÓN": "descripcion",
    "ESTADO": "estado", "% AVANCE": "avance", "RESPONSABLE": "responsable",
    "FECHA ACTUALIZACIÓN": "fecha", "OBSERVACIONES": "observaciones",
}


def normalizar(tipo, valor):
    """'PAQUETE 9', 'Paquete 9 ', 'COMPONENTE4' → 'Paquete 9', 'Componente 4'."""
    m = re.match(rf"\s*{tipo}\s*([\d.]+)", str(valor or ""), re.I)
    return f"{tipo.title()} {m.group(1)}" if m else str(valor or "").strip()


def leer(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["CONTROL GENERAL"] if "CONTROL GENERAL" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    i_cab = next(i for i, f in enumerate(filas) if f and "COMPONENTE" in [str(c).strip() for c in f if c])
    cab = [str(c).strip() if c else "" for c in filas[i_cab]]
    idx = {COLUMNAS[c]: i for i, c in enumerate(cab) if c in COLUMNAS and COLUMNAS[c] not in ()}
    out, sector, paquete = [], "", ""
    for f in filas[i_cab + 1:]:
        comp = f[idx["componente"]] if f else None
        if not comp:
            continue
        sector = str(f[idx["sector"]]).strip().title() if f[idx["sector"]] else sector
        paquete = normalizar("PAQUETE", f[idx["paquete"]]) if f[idx["paquete"]] else paquete
        fila = {k: f[i] for k, i in idx.items()}
        fecha = fila.get("fecha")
        fila.update({
            "sector": sector,
            "paquete": paquete,
            "componente": normalizar("COMPONENTE", comp),
            "descripcion": str(fila.get("descripcion") or "").strip(),
            "estado": str(fila.get("estado") or "Pendiente").strip(),
            "avance": fila.get("avance") if fila.get("avance") is not None else "",
            "responsable": str(fila.get("responsable") or "").strip(),
            "fecha": fecha.strftime("%Y-%m-%d") if isinstance(fecha, (date, datetime)) else str(fecha or ""),
            "observaciones": str(fila.get("observaciones") or "").strip(),
        })
        out.append(fila)
    return out


def main():
    filas = leer(sys.argv[1])
    salida = Path(sys.argv[2])
    salida.mkdir(parents=True, exist_ok=True)
    campos = ["sector", "paquete", "componente", "descripcion", "estado", "avance", "responsable", "fecha", "observaciones"]
    with open(salida / "componentes.csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos)
        w.writeheader()
        w.writerows(filas)
    print(f"{len(filas)} componentes en {len({f['paquete'] for f in filas})} paquetes y {len({f['sector'] for f in filas})} sectores")


if __name__ == "__main__":
    main()
